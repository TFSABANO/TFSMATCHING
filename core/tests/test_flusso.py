"""Test end-to-end: import dei file di esempio, abbinamento automatico, viste, Aruba (simulata)."""
import base64
import json
from datetime import date
from decimal import Decimal as D

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from core import matching, servizi
from core.matching import numero_in_descrizione
from core.models import Abbinamento, Azienda, ContoBancario, CredenzialiAruba, Fattura, Membro, MovimentoBancario

from . import campioni

ALFA = ("Cliente Alfa Srl", "11111111111")
BETA = ("Beta Consulting Sas", "22222222222")
NOI = ("La Mia Azienda Srl", campioni.PIVA_AZIENDA)


def fatture_esempio():
    def attiva(numero, data, cliente, totale, imponibile):
        imposta = f"{D(totale) - D(imponibile):.2f}"
        return campioni.fattura_xml(numero, data, cedente=NOI, cessionario=cliente, totale=totale,
                                    imponibile=imponibile, imposta=imposta)

    return campioni.zip_di({
        "f15.xml": attiva("2026/15", "2026-01-15", ALFA, "1220.00", "1000.00"),
        "f16.xml": attiva("2026/16", "2026-01-20", BETA, "900.00", "737.70"),
        "f17.xml.p7m": campioni.firma_p7m(attiva("2026/17", "2026-01-25", BETA, "600.00", "491.80")),
        "f18.xml": attiva("2026/18", "2026-02-01", ALFA, "305.00", "250.00"),
        "p77.xml": campioni.fattura_xml("A-77", "2026-01-20", totale="610.00", imponibile="500.00", imposta="110.00",
                                        scadenza="2026-02-03"),
        "tel.xml": campioni.fattura_xml("TI-998877", "2026-01-31", cedente=("TELECOM ITALIA S.p.A.", "00488410010"),
                                        totale="49.90", imponibile="40.90", imposta="9.00"),
    })


class FlussoCompletoTest(TestCase):
    def setUp(self):
        self.az = Azienda.objects.create(ragione_sociale="La Mia Azienda Srl", partita_iva="IT" + campioni.PIVA_AZIENDA)
        self.revolut = ContoBancario.objects.create(azienda=self.az, nome="Revolut", banca="revolut")
        self.sumup = ContoBancario.objects.create(azienda=self.az, nome="SumUp", banca="sumup")
        self.vivid = ContoBancario.objects.create(azienda=self.az, nome="Vivid", banca="vivid")

    def importa_tutto(self):
        imp = servizi.importa_fatture_file(self.az, "export_aruba.zip", fatture_esempio())
        self.assertEqual(imp.note, "")
        self.assertEqual(imp.righe_importate, 6)
        servizi.importa_corrispettivi(self.az, "corr.csv", campioni.CSV_CORRISPETTIVI_LUNGO.encode())
        servizi.importa_estratto(self.az, self.revolut, "rev.pdf", campioni.pdf_revolut())
        servizi.importa_estratto(self.az, self.sumup, "sumup.pdf", campioni.pdf_sumup())
        servizi.importa_estratto(self.az, self.vivid, "vivid.pdf", campioni.pdf_vivid())

    def proposte_per(self, importo):
        mov = MovimentoBancario.objects.get(azienda=self.az, importo=importo)
        return list(mov.abbinamenti.filter(stato=Abbinamento.PROPOSTO))

    def test_import_idempotente(self):
        self.importa_tutto()
        n_mov = MovimentoBancario.objects.count()
        imp = servizi.importa_estratto(self.az, self.revolut, "rev.pdf", campioni.pdf_revolut())
        self.assertEqual((imp.righe_importate, imp.righe_saltate), (0, 4))
        self.assertEqual(MovimentoBancario.objects.count(), n_mov)
        imp = servizi.importa_fatture_file(self.az, "export_aruba.zip", fatture_esempio())
        self.assertEqual((imp.righe_importate, imp.righe_saltate), (0, 6))

    def test_direzione_fatture(self):
        self.importa_tutto()
        self.assertEqual(Fattura.objects.filter(direzione="attiva").count(), 4)
        p = Fattura.objects.get(numero="A-77")
        self.assertEqual((p.direzione, p.controparte), ("passiva", "Fornitore Srl"))

    def test_abbinamenti_automatici(self):
        self.importa_tutto()
        matching.genera_proposte(self.az)

        [a] = self.proposte_per(D("1220.00"))
        self.assertEqual(a.fattura.numero, "2026/15")
        self.assertEqual(a.punteggio, 100)

        [a] = self.proposte_per(D("-610.00"))
        self.assertEqual(a.fattura.numero, "A-77")

        self.assertEqual(self.proposte_per(D("-35.99")), [])  # spesa con carta: nessuna fattura

        prop = self.proposte_per(D("1500.00"))  # bonifico cumulativo
        self.assertEqual(sorted(a.fattura.numero for a in prop), ["2026/16", "2026/17"])
        self.assertTrue(prop[0].gruppo and prop[0].gruppo == prop[1].gruppo)

        [a] = self.proposte_per(D("305.00"))
        self.assertEqual(a.fattura.numero, "2026/18")
        [a] = self.proposte_per(D("-49.90"))
        self.assertEqual(a.fattura.numero, "TI-998877")

        # SumUp: accredito netto di due giorni di incassi POS (01/02 300 + 02/02 195 = 495 lordi)
        prop = self.proposte_per(D("486.35"))
        self.assertEqual(sorted((a.corrispettivo.data.day, a.importo_documento) for a in prop), [(1, D("300.00")), (2, D("195.00"))])
        self.assertEqual(sum(a.importo for a in prop), D("486.35"))
        prop = self.proposte_per(D("688.04"))
        self.assertEqual(sorted(a.corrispettivo.data.day for a in prop), [3, 4])
        self.assertEqual(self.proposte_per(D("-1.50")), [])

    def test_conferma_rifiuto_e_residui(self):
        self.importa_tutto()
        matching.genera_proposte(self.az)
        prop = self.proposte_per(D("1500.00"))
        matching.conferma(prop[0])  # conferma tutto il gruppo
        self.assertEqual(Abbinamento.objects.filter(gruppo=prop[0].gruppo, stato="confermato").count(), 2)
        self.assertEqual(Fattura.objects.get(numero="2026/16").residuo, D("0"))
        self.assertEqual(MovimentoBancario.objects.get(importo=D("1500.00")).residuo, D("0"))

        [a] = self.proposte_per(D("1220.00"))
        matching.rifiuta(a)
        matching.genera_proposte(self.az)
        self.assertEqual(self.proposte_per(D("1220.00")), [])  # non viene riproposta
        # le confermate non vengono toccate dal ricalcolo
        self.assertEqual(Abbinamento.objects.filter(stato="confermato").count(), 2)

    def test_conferma_sopra_soglia(self):
        self.importa_tutto()
        matching.genera_proposte(self.az)
        n = matching.conferma_sopra_soglia(self.az, soglia=90)
        self.assertGreaterEqual(n, 3)
        self.assertEqual(Fattura.objects.get(numero="2026/15").residuo, D("0"))

    def test_export_excel(self):
        from openpyxl import load_workbook
        import io

        from core.export import esporta_riconciliazione

        self.importa_tutto()
        matching.genera_proposte(self.az)
        matching.conferma_sopra_soglia(self.az, soglia=0)
        wb = load_workbook(io.BytesIO(esporta_riconciliazione(self.az)))
        self.assertEqual(wb.sheetnames, ["Movimenti", "Fatture", "Corrispettivi"])
        fatture = {r[2]: r[10] for r in wb["Fatture"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(fatture["2026/15"], "Abbinato")
        corr = list(wb["Corrispettivi"].iter_rows(min_row=2, values_only=True))
        commissioni = sum(r[5] for r in corr)
        self.assertAlmostEqual(commissioni, 495 - 486.35 + 700 - 688.04, places=2)


class NumeroFatturaTest(TestCase):
    def test_varianti(self):
        self.assertTrue(numero_in_descrizione("2026/15", "Saldo fattura 2026/15"))
        self.assertTrue(numero_in_descrizione("2026/15", "pagamento ft 15/2026"))
        self.assertTrue(numero_in_descrizione("15", "pagamento fatt. n. 15 del 3/1"))
        self.assertTrue(numero_in_descrizione("FPR 7/26", "fattura n 7"))
        self.assertFalse(numero_in_descrizione("15", "bonifico 150 euro"))
        self.assertFalse(numero_in_descrizione("15", "iban 15"))
        self.assertFalse(numero_in_descrizione("2026/15", "fattura 2026/150"))


class FakeResp:
    def __init__(self, status, dati):
        self.status_code, self._dati, self.text = status, dati, json.dumps(dati)

    def json(self):
        return self._dati


class FakeSession:
    """Simula le API Aruba: una fattura emessa e una ricevuta."""

    def __init__(self, files):
        self.files = files
        self.chiamate = []

    def post(self, url, data=None, headers=None, timeout=None):
        self.chiamate.append(("POST", url))
        if data["password"] != "segreta":
            return FakeResp(401, {"error": "invalid_grant"})
        return FakeResp(200, {"access_token": "tok", "expires_in": 1800})

    def get(self, url, params=None, headers=None, timeout=None):
        self.chiamate.append(("GET", url, params))
        assert headers["Authorization"] == "Bearer tok"
        direzione = "out" if "/invoice/out/" in url else "in"
        if url.endswith("findByUsername"):
            nomi = [n for n, (d, _) in self.files.items() if d == direzione]
            return FakeResp(200, {"content": [{"filename": n, "idSdi": "123"} for n in nomi], "totalPages": 1})
        nome = params["filename"]
        return FakeResp(200, {"filename": nome, "file": base64.b64encode(self.files[nome][1]).decode()})


class ArubaTest(TestCase):
    def setUp(self):
        self.az = Azienda.objects.create(ragione_sociale="La Mia Azienda Srl", partita_iva=campioni.PIVA_AZIENDA)
        cred = CredenzialiAruba(azienda=self.az, username="utente", ambiente="demo")
        cred.password = "segreta"
        cred.save()

    def test_password_cifrata(self):
        cred = CredenzialiAruba.objects.get(azienda=self.az)
        self.assertNotIn("segreta", cred._password)
        self.assertEqual(cred.password, "segreta")

    def test_sincronizzazione(self):
        from core.aruba import ArubaClient

        xml_out = campioni.fattura_xml("2026/40", "2026-03-01", cedente=NOI, cessionario=ALFA)
        xml_in = campioni.fattura_xml("F-1", "2026-03-02")
        sess = FakeSession({
            "IT01234567890_aaa01.xml.p7m": ("out", campioni.firma_p7m(xml_out)),
            "IT09876543210_bbb02.xml": ("in", xml_in),
        })
        client = ArubaClient("utente", "segreta", "demo", session=sess)
        imp = servizi.sincronizza_aruba(self.az, date(2026, 1, 1), date(2026, 3, 31), client=client)
        self.assertEqual(imp.note, "")
        self.assertEqual(imp.righe_importate, 2)
        self.assertEqual(Fattura.objects.get(numero="2026/40").direzione, "attiva")
        self.assertEqual(Fattura.objects.get(numero="F-1").direzione, "passiva")
        self.assertTrue(any("demows.fatturazioneelettronica.aruba.it" in c[1] for c in sess.chiamate if c[0] == "GET"))
        # seconda sincronizzazione: i file già scaricati non vengono richiesti di nuovo
        sess.chiamate.clear()
        servizi.sincronizza_aruba(self.az, date(2026, 1, 1), date(2026, 3, 31), client=client)
        self.assertFalse(any("getByFilename" in c[1] for c in sess.chiamate if c[0] == "GET"))


class VisteTest(TestCase):
    def setUp(self):
        self.az = Azienda.objects.create(ragione_sociale="La Mia Azienda Srl", partita_iva=campioni.PIVA_AZIENDA)
        self.altra = Azienda.objects.create(ragione_sociale="Altra Spa", partita_iva="99999999999")
        self.conto = ContoBancario.objects.create(azienda=self.az, nome="Revolut", banca="revolut")
        self.admin = User.objects.create_user("mario", password="x")
        Membro.objects.create(utente=self.admin, azienda=self.az, ruolo="admin")
        self.lettore = User.objects.create_user("luca", password="x")
        Membro.objects.create(utente=self.lettore, azienda=self.az, ruolo="lettura")

    def test_login_richiesto(self):
        r = self.client.get(reverse("dashboard"))
        self.assertEqual(r.status_code, 302)
        self.assertIn("/login/", r["Location"])

    def test_flusso_da_interfaccia(self):
        self.client.force_login(self.admin)
        r = self.client.post(reverse("importa"), {
            "azione": "fatture", "fatt-file": [SimpleUploadedFile("export.zip", fatture_esempio())], "fatt-direzione": "",
        })
        self.assertEqual(r.status_code, 302)
        self.client.post(reverse("importa"), {
            "azione": "estratto", "estr-conto": self.conto.pk,
            "estr-file": [SimpleUploadedFile("rev.pdf", campioni.pdf_revolut())],
        })
        self.client.post(reverse("importa"), {
            "azione": "corrispettivi", "corr-metodo_default": "carta",
            "corr-file": [SimpleUploadedFile("c.csv", campioni.CSV_CORRISPETTIVI_LUNGO.encode())],
        })
        self.assertEqual(MovimentoBancario.objects.filter(azienda=self.az).count(), 4)
        self.client.post(reverse("genera_proposte"))
        r = self.client.get(reverse("riconciliazione"))
        self.assertContains(r, "2026/15")
        a = Abbinamento.objects.filter(azienda=self.az, fattura__numero="2026/15").first()
        self.client.post(reverse("azione_abbinamento", args=[a.pk, "conferma"]))
        a.refresh_from_db()
        self.assertEqual(a.stato, "confermato")

        # abbinamento manuale della spesa Amazon a una fattura passiva (es. registrata a mano)
        mov = MovimentoBancario.objects.get(importo=D("-35.99"))
        r = self.client.get(reverse("movimento", args=[mov.pk]))
        self.assertEqual(r.status_code, 200)
        tel = Fattura.objects.get(numero="TI-998877")
        self.client.post(reverse("movimento", args=[mov.pk]), {"tipo": "f", "documento_id": tel.pk, "importo": "35.99"})
        self.assertEqual(mov.residuo, D("0"))

        for nome in ["dashboard", "fatture", "corrispettivi", "importazioni", "impostazioni", "importa"]:
            self.assertEqual(self.client.get(reverse(nome)).status_code, 200, nome)
        for stato in ["proposti", "aperti", "abbinati", "tutti"]:
            self.assertEqual(self.client.get(reverse("riconciliazione"), {"stato": stato}).status_code, 200)
        r = self.client.get(reverse("esporta"))
        self.assertEqual(r["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    def test_isolamento_aziende(self):
        conto_altra = ContoBancario.objects.create(azienda=self.altra, nome="X", banca="generico")
        mov = MovimentoBancario.objects.create(azienda=self.altra, conto=conto_altra, data=date(2026, 1, 1),
                                               descrizione="segreto", importo=D("10"), impronta="x")
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get(reverse("movimento", args=[mov.pk])).status_code, 404)
        self.assertNotContains(self.client.get(reverse("riconciliazione"), {"stato": "tutti"}), "segreto")
        # non può passare a un'azienda di cui non è membro
        self.client.post(reverse("cambia_azienda"), {"azienda_id": self.altra.pk})
        self.assertNotContains(self.client.get(reverse("riconciliazione"), {"stato": "tutti"}), "segreto")

    def test_ruolo_lettura(self):
        self.client.force_login(self.lettore)
        self.assertEqual(self.client.get(reverse("dashboard")).status_code, 200)
        self.assertEqual(self.client.get(reverse("importa")).status_code, 403)
        self.assertEqual(self.client.get(reverse("impostazioni")).status_code, 403)
        self.assertEqual(self.client.post(reverse("genera_proposte")).status_code, 403)
