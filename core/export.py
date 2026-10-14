"""Export della riconciliazione in Excel."""
import io
from collections import defaultdict
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Abbinamento, Corrispettivo, Fattura, MovimentoBancario

VERDE = PatternFill("solid", fgColor="D9F2D9")
GIALLO = PatternFill("solid", fgColor="FFF2CC")
ROSSO = PatternFill("solid", fgColor="F8D7DA")


def _foglio(wb, titolo, intestazioni, righe, colore=None):
    ws = wb.create_sheet(titolo)
    ws.append(intestazioni)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in righe:
        ws.append(r)
        if colore:
            fill = colore(r)
            if fill:
                for c in ws[ws.max_row]:
                    c.fill = fill
    for i, h in enumerate(intestazioni, start=1):
        larghezza = max([len(str(h))] + [len(str(r[i - 1])) for r in righe[:500] if r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, larghezza + 2), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def _stato(residuo, totale):
    if residuo <= Decimal("0.01"):
        return "Abbinato"
    if residuo < totale:
        return "Parziale"
    return "Da abbinare"


def _colore_stato(indice):
    return lambda r: {"Abbinato": VERDE, "Parziale": GIALLO, "Da abbinare": ROSSO}.get(r[indice])


def esporta_riconciliazione(azienda, dal=None, al=None) -> bytes:
    conf = Abbinamento.objects.filter(azienda=azienda, stato=Abbinamento.CONFERMATO).select_related("fattura", "corrispettivo")
    per_mov, per_fatt, per_corr = defaultdict(list), defaultdict(list), defaultdict(list)
    for a in conf:
        per_mov[a.movimento_id].append(a)
        if a.fattura_id:
            per_fatt[a.fattura_id].append(a)
        else:
            per_corr[a.corrispettivo_id].append(a)

    def filtra(qs):
        if dal:
            qs = qs.filter(data__gte=dal)
        if al:
            qs = qs.filter(data__lte=al)
        return qs

    wb = Workbook()
    wb.remove(wb.active)

    righe = []
    for m in filtra(MovimentoBancario.objects.filter(azienda=azienda).select_related("conto")).order_by("data", "id"):
        abb = per_mov.get(m.pk, [])
        coperto = sum((a.importo for a in abb), Decimal("0"))
        residuo = abs(m.importo) - coperto
        documenti = "; ".join(
            f"Ft {a.fattura.numero} {a.fattura.controparte}" if a.fattura_id else f"Corr. {a.corrispettivo.data:%d/%m/%Y} {a.corrispettivo.get_metodo_display()}"
            for a in abb
        )
        righe.append([m.data, m.conto.nome, m.descrizione, float(m.importo), float(residuo), _stato(residuo, abs(m.importo)), documenti])
    _foglio(wb, "Movimenti", ["Data", "Conto", "Descrizione", "Importo", "Residuo", "Stato", "Documenti abbinati"], righe, _colore_stato(5))

    righe = []
    for f in filtra(Fattura.objects.filter(azienda=azienda)).order_by("data", "numero"):
        abb = per_fatt.get(f.pk, [])
        coperto = sum((a.importo_documento for a in abb), Decimal("0"))
        residuo = f.importo_da_abbinare - coperto
        pagamenti = "; ".join(f"{a.movimento.data:%d/%m/%Y} {a.importo}" for a in abb)
        righe.append([
            f.get_direzione_display(), f.tipo_documento, f.numero, f.data, f.controparte, f.controparte_piva,
            float(f.totale), float(f.netto_a_pagare), f.data_scadenza, float(residuo), _stato(residuo, f.importo_da_abbinare), pagamenti,
        ])
    _foglio(wb, "Fatture", ["Tipo", "TD", "Numero", "Data", "Controparte", "P.IVA/CF", "Totale", "Netto a pagare",
                            "Scadenza", "Residuo", "Stato", "Pagamenti"], righe, _colore_stato(10))

    righe = []
    for c in filtra(Corrispettivo.objects.filter(azienda=azienda)).order_by("data", "metodo"):
        abb = per_corr.get(c.pk, [])
        coperto = sum((a.importo_documento for a in abb), Decimal("0"))
        accreditato = sum((a.importo for a in abb), Decimal("0"))
        residuo = c.importo - coperto
        stato = "Contanti" if c.metodo == "contanti" and not abb else _stato(residuo, c.importo)
        righe.append([c.data, c.get_metodo_display(), c.riferimento, float(c.importo), float(accreditato),
                      float(coperto - accreditato), float(residuo), stato])
    _foglio(wb, "Corrispettivi", ["Data", "Metodo", "Riferimento", "Lordo", "Accreditato", "Commissioni", "Residuo", "Stato"],
            righe, _colore_stato(7))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
